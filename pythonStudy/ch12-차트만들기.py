import openpyxl
from openpyxl.chart import BarChart, Reference, LineChart

company_list = ["애플", "아마존", "테슬라"]
for company in company_list:
    book = openpyxl.load_workbook(f"./엑셀데이터/2024년_광고비_{company}.xlsx")
    sheet = book.active
    # 월별 total 광고비로 막대 차트 만들기
    bar_chart = BarChart()
    bar_chart.title = f"{company} 월별 광고비"
    data = Reference(sheet, range_string="Sheet1!G1:G13")
    bar_chart.add_data(data, titles_from_data=True)
    sheet.add_chart(bar_chart, "H1")

    # 월별 매체별 광고비로 꺾은선 차트 만들기
    line_chart = LineChart()
    line_chart.title = f"{company} 월별 광고비"
    for column in ["C", "D", "E", "F"]:
        data = Reference(sheet, range_string=f"Sheet1!{column}1:{column}13")
        line_chart.add_data(data, titles_from_data=True)
    sheet.add_chart(line_chart, "H14")

    book.save(f"./엑셀데이터/차트만들기_{company}.xlsx")